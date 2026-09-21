from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from core.types import ChunkLabel, Domain, SourceStatus
from corpus.api import get_chunk, get_corpus_version, search, supported_domains
from corpus import extract_doc
from corpus.chunker import chunk_document
from corpus.extract_doc import extract_document, normalize_pages
from corpus.intake import (
    SourceRecheckResult,
    ingest_file,
    ingest_text,
    ingest_url,
    recheck_url_sources,
)
from corpus.metadata import METADATA_SCHEMA, propose_metadata
from corpus.store import (
    ChunkRecord,
    CorpusVersionRecord,
    SourceRecord,
    bump_corpus_version,
    compute_corpus_version,
    create_chunk,
    create_corpus_version,
    create_source,
    delete_chunk,
    delete_corpus_version,
    delete_source,
    get_chunk_record,
    get_corpus_version_record,
    get_current_corpus_version,
    get_source,
    list_sources,
    update_chunk,
    update_corpus_version,
    update_source,
)
from infra.audit import recent_events


def test_stub_corpus_api_exposes_contract_chunks() -> None:
    domains = supported_domains()
    chunks = [chunk for domain in domains for chunk in search("", [domain])]

    assert domains == [Domain.CONDUCT_SCORE, Domain.COURSE_WITHDRAWAL, Domain.GRADE_APPEAL]
    assert get_corpus_version() == "stub-v1"
    assert len(chunks) == 12
    assert sum(chunk.label is ChunkLabel.HUMAN_ONLY for chunk in chunks) == 2
    assert sum(chunk.transitional_clause for chunk in chunks) == 1
    assert get_chunk(chunks[0].chunk_id) == chunks[0]


def test_search_filters_domains_and_limits_results() -> None:
    chunks = search(
        "Hạn chót rút học phần là khi nào?",
        [Domain.COURSE_WITHDRAWAL],
        top_k=2,
        at=datetime(2026, 9, 21, tzinfo=timezone.utc),
    )

    assert len(chunks) == 2
    assert {chunk.domain for chunk in chunks} == {Domain.COURSE_WITHDRAWAL}


def test_store_versions_only_active_sources(tmp_path: Path) -> None:
    database_path = tmp_path / "corpus.db"
    pending = SourceRecord(doc_id="source-1", content_hash="first")

    stored_pending = create_source(pending, database_path=database_path)
    empty_version = compute_corpus_version(database_path=database_path)
    assert get_source("source-1", database_path=database_path) == stored_pending

    update_source(
        replace(stored_pending, status=SourceStatus.ACTIVE),
        database_path=database_path,
    )
    active_version = bump_corpus_version(
        "ADMIN:tester", "Kích hoạt nguồn", database_path=database_path
    )
    assert active_version != empty_version
    assert get_current_corpus_version(database_path=database_path) == active_version

    update_source(
        replace(stored_pending, status=SourceStatus.SUPERSEDED),
        database_path=database_path,
    )
    assert (
        bump_corpus_version("ADMIN:tester", "Hạ cấp nguồn", database_path=database_path)
        == empty_version
    )


def test_store_crud_for_chunks_and_versions(tmp_path: Path) -> None:
    database_path = tmp_path / "corpus.db"
    source = SourceRecord(doc_id="source-1", content_hash="first")
    chunk = ChunkRecord(
        chunk_id="chunk-1",
        doc_id=source.doc_id,
        breadcrumb="QĐ 1 · Điều 1",
        text="Nội dung nguồn.",
        domain=Domain.CONDUCT_SCORE,
    )
    version = CorpusVersionRecord("cv_manual", "ADMIN:tester", "Ghi chú", (source.doc_id,))

    create_source(source, database_path=database_path)
    create_chunk(chunk, database_path=database_path)
    assert get_chunk_record(chunk.chunk_id, database_path=database_path) == chunk

    updated_chunk = replace(chunk, label=ChunkLabel.AUTO_ANSWERABLE)
    update_chunk(updated_chunk, database_path=database_path)
    assert get_chunk_record(chunk.chunk_id, database_path=database_path) == updated_chunk

    stored_version = create_corpus_version(version, database_path=database_path)
    updated_version = replace(stored_version, note="Đã sửa")
    update_corpus_version(updated_version, database_path=database_path)
    assert (
        get_corpus_version_record(version.corpus_version, database_path=database_path)
        == updated_version
    )

    delete_chunk(chunk.chunk_id, database_path=database_path)
    delete_corpus_version(version.corpus_version, database_path=database_path)
    delete_source(source.doc_id, database_path=database_path)
    assert get_chunk_record(chunk.chunk_id, database_path=database_path) is None
    assert get_corpus_version_record(version.corpus_version, database_path=database_path) is None
    assert get_source(source.doc_id, database_path=database_path) is None


def test_intake_deduplicates_content_and_writes_audit(tmp_path: Path) -> None:
    database_path = tmp_path / "corpus.db"

    first = ingest_text(
        "Điều 1. Nội dung quy định.",
        actor="ADMIN:tester",
        database_path=database_path,
    )
    duplicate = ingest_text(
        "Điều 1. Nội dung quy định.",
        actor="ADMIN:tester",
        database_path=database_path,
    )

    assert first.created is True
    assert duplicate.created is False
    assert duplicate.message == "Tài liệu không thay đổi"
    assert len(list_sources(database_path=database_path)) == 1
    assert [event.action for event in recent_events(database_path=str(database_path))] == [
        "SOURCE_UPLOADED"
    ]


def test_intake_accepts_supported_files_and_manual_url(tmp_path: Path) -> None:
    database_path = tmp_path / "corpus.db"

    uploaded = ingest_file(
        b"%PDF-1.7", "quy_dinh.pdf", actor="ADMIN:tester", database_path=database_path
    )
    downloaded = ingest_url(
        "https://example.edu/quy-dinh.docx",
        actor="ADMIN:tester",
        fetch=lambda _: b"docx bytes",
        database_path=database_path,
    )

    assert uploaded.source.source_kind == "pdf"
    assert downloaded.source.source_kind == "url"
    assert downloaded.source.source_url == "https://example.edu/quy-dinh.docx"


def test_recheck_url_sources_reports_changes_and_creates_pending_source(tmp_path: Path) -> None:
    database_path = tmp_path / "corpus.db"
    original = ingest_url(
        "https://example.edu/quy-dinh.docx",
        actor="ADMIN:tester",
        fetch=lambda _: b"ban dau",
        database_path=database_path,
    ).source

    unchanged = recheck_url_sources(
        actor="ADMIN:tester", fetch=lambda _: b"ban dau", database_path=database_path
    )
    changed = recheck_url_sources(
        actor="ADMIN:tester", fetch=lambda _: b"ban moi", database_path=database_path
    )

    assert unchanged == [SourceRecheckResult(original, False, "Không đổi")]
    assert changed[0].changed is True
    assert changed[0].proposed_source is not None
    assert changed[0].proposed_source.status is SourceStatus.PENDING_REVIEW
    assert len(list_sources(database_path=database_path)) == 2
    assert [event.action for event in recent_events(database_path=str(database_path))].count(
        "SOURCE_RECHECKED"
    ) == 2


def test_metadata_proposal_uses_excerpt_and_keeps_unknown_fields_null(monkeypatch) -> None:
    documents = [f"TÀI LIỆU SỐ {number}\nĐiều {number}. Nội dung." for number in range(1, 7)]
    prompts: list[str] = []

    def fake_call(prompt, *, schema, step, case_id, temperature):
        prompts.append(prompt)
        number = next(
            (number for number in range(1, 7) if f"TÀI LIỆU SỐ {number}" in prompt), 1
        )
        assert schema == METADATA_SCHEMA
        assert step == "K3_metadata"
        assert temperature == 0.0
        return SimpleNamespace(
            ok=True,
            data={
                "document_id": f"DOC-{number}",
                "title": f"Tài liệu {number}",
                "issuer": None,
                "published_at": None,
                "effective_from": None,
                "effective_to": None,
                "applies_to": None,
                "cohorts": None,
                "supersedes": ["DOC-OLD"] if number == 1 else None,
                "transitional_clause": number == 1,
                "domains": ["conduct_score"],
                "status": "PENDING_REVIEW",
                "content_hash": None,
            },
            error=None,
        )

    monkeypatch.setattr("corpus.metadata.call_json", fake_call)
    proposals = [propose_metadata(document, case_id=f"case-{number}") for number, document in enumerate(documents, 1)]

    assert all(proposal.error is None and proposal.draft is not None for proposal in proposals)
    assert proposals[0].draft.transitional_clause is True
    assert proposals[0].draft.supersedes == ("DOC-OLD",)
    assert proposals[1].draft.transitional_clause is False
    assert proposals[1].draft.issuer is None
    assert proposals[1].draft.cohorts is None
    assert propose_metadata("x" * 3_001, case_id="case-long").error is None
    assert prompts[-1].endswith("x" * 3_000)


def test_chunker_keeps_legal_units_breadcrumbs_and_long_clause_content() -> None:
    first = chunk_document(
        "Điều 8. Rút học phần\nKhoản 2. Sinh viên nộp đơn.\nĐiểm a) Nộp trước hạn.",
        doc_id="QD-3150",
        document_title="QĐ 3150/2026",
        domain=Domain.COURSE_WITHDRAWAL,
    )
    second = chunk_document(
        "Điều 4. Phúc khảo\nKhoản 1. Chuyên viên tiếp nhận hồ sơ.",
        doc_id="QD-4000",
        document_title="QĐ 4000/2026",
        domain=Domain.GRADE_APPEAL,
    )
    long_clause = " ".join(["nội_dung"] * 801)
    third = chunk_document(
        f"Điều 3. Điểm rèn luyện\nKhoản 2. {long_clause}",
        doc_id="QD-5000",
        document_title="QĐ 5000/2026",
        domain=Domain.CONDUCT_SCORE,
    )

    assert all(chunk.text.strip() for chunk in first + second + third)
    assert "QĐ 3150/2026 · Điều 8 · Khoản 2" in {
        chunk.breadcrumb for chunk in first
    }
    assert "QĐ 3150/2026 · Điều 8 · Khoản 2 · Điểm a" in {
        chunk.breadcrumb for chunk in first
    }
    assert "Điểm a) Nộp trước hạn." in "\n".join(chunk.text for chunk in first)
    assert "Chuyên viên tiếp nhận hồ sơ." in "\n".join(chunk.text for chunk in second)
    long_chunks = [chunk for chunk in third if chunk.clause_no == "2"]
    assert len(long_chunks) == 2
    assert {chunk.breadcrumb for chunk in long_chunks} == {"QĐ 5000/2026 · Điều 3 · Khoản 2"}
    assert all(chunk.token_count <= 800 for chunk in long_chunks)
    assert [chunk.ordinal for chunk in first] == list(range(1, len(first) + 1))


def test_normalize_pages_removes_repeated_margins_and_keeps_legal_headings() -> None:
    text = normalize_pages(
        [
            "TRƯỜNG ĐẠI HỌC\nĐiều 1. Phạm vi áp dụng\nSinh viên thực hiện\ntheo quy định.\nTrang 1",
            "TRƯỜNG ĐẠI HỌC\nĐiều 2. Đối tượng áp dụng\nKhoản 1. Thời hạn nộp\nHồ sơ được tiếp nhận\ntrong giờ hành chính.\nTrang 2",
            "TRƯỜNG ĐẠI HỌC\nĐiều 3. Hồ sơ\nĐiểm a) Hồ sơ cần có\nĐơn đề nghị hợp lệ.\nTrang 3",
            "TRƯỜNG ĐẠI HỌC\nĐiều 4. Trình tự\nNộp hồ sơ tại phòng CTSV.\nTrang 4",
            "TRƯỜNG ĐẠI HỌC\nĐiều 5. Thời gian\nGiải quyết trong năm ngày.\nTrang 5",
            "TRƯỜNG ĐẠI HỌC\nĐiều 6. Thi hành\nQuy định này có hiệu lực.\nTrang 6",
        ]
    )

    assert "TRƯỜNG ĐẠI HỌC" not in text
    assert "Trang 1" in text
    assert "Điều 1. Phạm vi áp dụng" in text
    assert all(f"Điều {number}." in text for number in range(1, 7))
    assert "Khoản 1. Thời hạn nộp" in text
    assert "Điểm a) Hồ sơ cần có" in text
    assert "Sinh viên thực hiện theo quy định." in text


def test_extract_document_reads_docx_and_pdf_pages(monkeypatch) -> None:
    document = Document()
    document.add_paragraph("Điều 2. Điều kiện")
    document.add_paragraph("Sinh viên nộp đơn hợp lệ.")
    content = BytesIO()
    document.save(content)

    assert "Điều 2. Điều kiện" in extract_document(content.getvalue(), "quy_dinh.docx")

    class FakePdf:
        pages = [SimpleNamespace(extract_text=lambda: "Điều 3. Thủ tục")]

        def __enter__(self):
            return self

        def __exit__(self, *_):
            return None

    monkeypatch.setattr(
        extract_doc, "pdfplumber", SimpleNamespace(open=lambda _: FakePdf())
    )
    assert extract_document(b"%PDF", "quy_dinh.pdf") == "Điều 3. Thủ tục"
